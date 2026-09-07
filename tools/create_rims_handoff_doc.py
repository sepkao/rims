from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = r"C:\Users\User\rims\RIMS_Work_Summary_Before_Git_Merge_2026-09-04.docx"
FONT = "Leelawadee UI"
BROWN = "3B241D"
ACCENT = "B86F52"
PALE = "F7F1ED"
PALE_ALT = "FBF8F5"
GRAY = "666666"
BORDER = "D9D9D9"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_run_font(run, name=FONT, size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph(p, before=0, after=6, line=1.15):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line


def add_body(doc, text, bold_lead=None, after=7):
    p = doc.add_paragraph()
    style_paragraph(p, after=after, line=1.2)
    if bold_lead and text.startswith(bold_lead):
        a = p.add_run(bold_lead)
        set_run_font(a, bold=True, size=11)
        b = p.add_run(text[len(bold_lead):])
        set_run_font(b, size=11)
    else:
        r = p.add_run(text)
        set_run_font(r, size=11)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    style_paragraph(p, after=3, line=1.15)
    for r in p.runs:
        set_run_font(r, size=10.5)
    if not p.runs:
        set_run_font(p.add_run(text), size=10.5)
    else:
        p.runs[0].text = text
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    style_paragraph(p, after=4, line=1.15)
    if p.runs:
        p.runs[0].text = text
        set_run_font(p.runs[0], size=10.5)
    else:
        set_run_font(p.add_run(text), size=10.5)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(10 if level == 1 else 7)
    p.paragraph_format.space_after = Pt(5)
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, BROWN)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        if widths:
            cell.width = Inches(widths[i])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(p, after=0, line=1.0)
        set_run_font(p.add_run(text), size=9.5, bold=True, color="FFFFFF")
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cell = cells[i]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if widths:
                cell.width = Inches(widths[i])
            if ridx % 2:
                set_cell_shading(cell, PALE_ALT)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            style_paragraph(p, after=0, line=1.08)
            set_run_font(p.add_run(str(text)), size=9.2)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("RIMS  สรุปงานก่อน Push และ Merge  |  4 กันยายน 2569")
    set_run_font(r, size=8, color=GRAY)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)
add_footer(section)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = FONT
normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal.font.size = Pt(11)

title = styles["Title"]
title.font.name = FONT
title._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
title.font.size = Pt(28)
title.font.bold = True
title.font.color.rgb = RGBColor(0, 0, 0)
title_ppr = title._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)
for level, size in ((1, 17), (2, 13)):
    st = styles[f"Heading {level}"]
    st.font.name = FONT
    st._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0, 0, 0)

# Cover
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(76)
p.paragraph_format.space_after = Pt(12)
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = p.add_run("SHABU RIMS")
set_run_font(r, size=12, bold=True, color=ACCENT)

p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
p.paragraph_format.space_after = Pt(12)
p.add_run("สรุปการพัฒนาและปรับปรุงระบบ RIMS")

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(30)
r = p.add_run("เอกสารส่งต่องานก่อน Push และ Merge เข้า Git")
set_run_font(r, size=15, color=GRAY)

add_body(doc, "เอกสารฉบับนี้สรุปการแก้ปัญหาการเข้าถึงระบบผ่าน QR การเปิดใช้งานชั่วคราวผ่าน Cloudflare Tunnel ฟีเจอร์แจ้งเตือนสต็อก Prep การรวมหน้าตั้งค่าวัตถุดิบ และผลการตรวจสอบความพร้อมของโครงการ เพื่อให้ทีมเข้าใจผลกระทบของการเปลี่ยนแปลงและตรวจสอบก่อนนำโค้ดเข้าสาขาหลัก")

meta = doc.add_table(rows=4, cols=2)
meta.alignment = WD_TABLE_ALIGNMENT.LEFT
meta.autofit = False
set_table_borders(meta)
for i, (label, value) in enumerate([
    ("วันที่จัดทำ", "4 กันยายน 2569"),
    ("ผู้รับเอกสาร", "ทีมพัฒนาและผู้ตรวจสอบโค้ด"),
    ("สถานะ", "พร้อมสำหรับการตรวจสอบก่อน Merge"),
    ("ขอบเขต", "API  Customer App  Internal App  Supabase Migration"),
]):
    meta.rows[i].cells[0].width = Inches(1.55)
    meta.rows[i].cells[1].width = Inches(5.2)
    for cell in meta.rows[i].cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, 140, 140, 140, 140)
    set_cell_shading(meta.rows[i].cells[0], PALE)
    p0 = meta.rows[i].cells[0].paragraphs[0]
    style_paragraph(p0, after=0)
    set_run_font(p0.add_run(label), size=10, bold=True)
    p1 = meta.rows[i].cells[1].paragraphs[0]
    style_paragraph(p1, after=0)
    set_run_font(p1.add_run(value), size=10)

doc.add_page_break()

add_heading(doc, "1 บทสรุปสำหรับทีม", 1)
add_body(doc, "งานรอบนี้เพิ่มความสามารถหลักด้านการควบคุมสต็อก Prep และการจัดการข้อมูลวัตถุดิบ พร้อมแก้การเชื่อมต่อ QR สำหรับการสาธิตข้ามเครือข่าย ระบบผ่านการ build ทั้งสามแอปและการทดสอบ API แล้ว แต่การใช้งานผ่าน Quick Tunnel ยังเป็นแบบชั่วคราว และมีรายการด้านสภาพแวดล้อมที่ควรปรับก่อน Production")

add_table(doc, ["หัวข้อ", "ผลลัพธ์", "สถานะ"], [
    ("QR และเครือข่าย", "ระบุสาเหตุจาก LAN IP เดิมหรือบริการไม่ตอบสนอง และเชื่อมผ่าน Quick Tunnel", "เสร็จสำหรับ Demo"),
    ("เกณฑ์ขั้นต่ำ Prep", "Owner ตั้งค่าเกณฑ์ขั้นต่ำ และ Staff Kitchen เห็นการแจ้งเตือนเมื่อสต็อกต่ำ", "พัฒนาแล้ว"),
    ("ตั้งค่าวัตถุดิบ", "แก้ชื่อ น้ำหนักต่อถาด เกณฑ์ขั้นต่ำ และเก็บถาวรในหน้าเดียว", "พัฒนาแล้ว"),
    ("ฐานข้อมูล", "เพิ่ม is_active และ index พร้อมนำ migration ไปใช้กับ Supabase", "นำไปใช้แล้ว"),
    ("การตรวจสอบ", "Build สำเร็จ และ API tests ผ่าน 15 จาก 15", "ผ่าน"),
    ("Production readiness", "ยังต้องจัดการ hosting secrets domain CORS cookies และ E2E", "ต้องดำเนินการต่อ"),
], widths=[1.65, 3.95, 1.25])

add_heading(doc, "สิ่งที่ทีมควรรู้ก่อน Merge", 2)
add_bullet(doc, "ไม่มีการลบวัตถุดิบจริง การเก็บถาวรใช้ฟิลด์ is_active เพื่อรักษาประวัติสต็อกและรายการสั่งซื้อ")
add_bullet(doc, "วัตถุดิบที่เก็บถาวรจะไม่ปรากฏในงานรับเข้า โอนสต็อก การแจ้งเตือน และการสร้างเมนูใหม่")
add_bullet(doc, "ระบบป้องกันการเก็บถาวรวัตถุดิบที่ยังถูกใช้ใน BOM ของเมนูที่เปิดใช้งาน")
add_bullet(doc, "นโยบายแจ้งเตือนการแปรรูปครอบคลุมวัตถุดิบประเภทเนื้อ ส่วนผักให้ Owner จัดซื้อเติมตามขอบเขตเดิม")
add_bullet(doc, "ไฟล์ .env จริงไม่ถูกติดตามด้วย Git และไม่มีค่าความลับบันทึกในเอกสารนี้")

doc.add_page_break()

add_heading(doc, "2 การแก้ปัญหา QR และการเข้าถึงหลายเครือข่าย", 1)
add_body(doc, "อาการ ERR CONNECTION TIMED OUT เกิดขึ้นเมื่อ QR ชี้ไปที่ 192.168.1.48:5173 ซึ่งเป็น IP ภายในเครือข่าย หากเปลี่ยน Wi-Fi เราเตอร์แจก IP ใหม่ หรือบริการ Vite ไม่ทำงาน อุปกรณ์ที่สแกน QR จะไม่สามารถเชื่อมต่อปลายทางเดิมได้")

add_heading(doc, "แนวทางที่ดำเนินการ", 2)
add_number(doc, "ใช้ Cloudflare Quick Tunnel เปิด API และ Customer App ให้เข้าถึงจากอินเทอร์เน็ตชั่วคราว")
add_number(doc, "กำหนด Customer App ให้เรียก API ผ่านตัวแปร VITE_API_URL และให้ค่าที่ระบุชนะค่า fallback ในโหมดพัฒนา")
add_number(doc, "เพิ่ม hostname ของ Customer Tunnel ใน Vite allowedHosts และเพิ่ม origin ใน CORS ของ API")
add_number(doc, "ตรวจสอบ Customer App และ API health ได้ HTTP 200 และตรวจสอบ CORS สำเร็จ")
add_number(doc, "QR เก่าที่ผูก token เดิมตอบกลับ 404 จึงต้องสร้าง QR ใหม่จาก URL และ token ปัจจุบัน")

add_heading(doc, "เส้นทางการเชื่อมต่อสำหรับ Demo", 2)
add_table(doc, ["ลำดับ", "การเชื่อมต่อ", "หน้าที่"], [
    ("1", "มือถือ → Customer Tunnel", "เปิดหน้า Customer App ผ่าน URL สาธารณะ"),
    ("2", "Customer App → API Tunnel", "ส่งคำขอข้อมูลและรายการสั่งซื้อไปยัง API"),
    ("3", "API → Supabase", "อ่านและบันทึกข้อมูลในฐานข้อมูล Cloud"),
], widths=[0.7, 2.6, 3.55])
add_body(doc, "ข้อจำกัดสำคัญ: Quick Tunnel เหมาะสำหรับการสาธิตเท่านั้น URL อาจเปลี่ยนเมื่อ restart และ process ของ Vite API และ cloudflared ต้องทำงานอยู่ตลอดเวลา URL ที่เคยสร้างจึงไม่ควรถูกใช้เป็น production endpoint")

add_heading(doc, "การใช้ Docker", 2)
add_body(doc, "Docker ไม่จำเป็นต่อขอบเขตปัจจุบัน เพราะ Customer และ Internal เป็น Vite static apps ส่วน API เป็น Hono บน Node และฐานข้อมูลอยู่บน Supabase Cloud อย่างไรก็ตาม Docker อาจเพิ่มภายหลังเพื่อทำให้สภาพแวดล้อม build และ deployment เหมือนกันทุกเครื่อง")

doc.add_page_break()

add_heading(doc, "3 ระบบเกณฑ์ขั้นต่ำ Prep และการแจ้งเตือน Staff Kitchen", 1)
add_body(doc, "Owner สามารถกำหนดจำนวนถาดขั้นต่ำที่ต้องมีใน Prep สำหรับวัตถุดิบแต่ละรายการ เมื่อจำนวนถาดปัจจุบันต่ำกว่าเกณฑ์ API จะคำนวณส่วนที่ขาดและแจ้ง Staff Kitchen ให้แปรรูปหรือโอนวัตถุดิบมาเติม")

add_table(doc, ["ส่วนงาน", "การเปลี่ยนแปลง"], [
    ("Owner", "แก้ค่า thaw_prep_threshold_plates ในหน้าตั้งค่าวัตถุดิบ"),
    ("API", "endpoint /staff/prep-alerts คำนวณจำนวนปัจจุบัน จำนวนที่ขาด กิโลกรัมแนะนำ และสต็อก Freezer"),
    ("Staff", "หน้ารายละเอียดแจ้งเตือน refresh ทุก 10 วินาทีและลิงก์ไปหน้าโอนเข้า Prep"),
    ("ทุกหน้าของ Staff", "คอมโพเนนต์แจ้งเตือนลอยแสดงเมื่อมีรายการต่ำกว่าเกณฑ์ และหายเองเมื่อเติมครบ"),
    ("หน้า Transfer", "รับ ingredient query parameter เพื่อเลือกวัตถุดิบจากการแจ้งเตือนล่วงหน้า"),
], widths=[1.55, 5.3])

add_heading(doc, "หลักการทำงาน", 2)
add_bullet(doc, "อ่านเฉพาะวัตถุดิบเนื้อที่เปิดใช้งาน")
add_bullet(doc, "เปรียบเทียบจำนวนถาดใน Prep กับเกณฑ์ขั้นต่ำของวัตถุดิบ")
add_bullet(doc, "คำนวณจำนวนถาดที่ขาดและน้ำหนักที่ควรนำมาแปรรูป")
add_bullet(doc, "แสดงข้อมูล Freezer เพื่อช่วย Staff ตัดสินใจว่ามีของเพียงพอหรือไม่")
add_bullet(doc, "เมื่อเติมสต็อกถึงเกณฑ์ การแจ้งเตือนจะหายโดยไม่สร้างข้อมูลแจ้งเตือนซ้ำในฐานข้อมูล")

add_heading(doc, "ขอบเขตทางธุรกิจ", 2)
add_body(doc, "การแจ้งเตือนให้ Staff Kitchen แปรรูปใช้กับเนื้อสัตว์ เพราะมีขั้นตอน Freezer ไป Prep และมีการคำนวณกิโลกรัม ส่วนผักใช้การติดตามจำนวนถาดใน Prep แต่การเติมเป็นความรับผิดชอบด้านการจัดซื้อของ Owner ตาม flow เดิม")

doc.add_page_break()

add_heading(doc, "4 หน้าตั้งค่าวัตถุดิบแบบรวม", 1)
add_body(doc, "ปรับหน้า Ingredient Settings ให้ Owner จัดการข้อมูลสำคัญในจุดเดียว ลดการแก้ข้อมูลโดยตรงในฐานข้อมูล และลดความเสี่ยงจากการลบรายการที่มีประวัติอ้างอิง")

add_table(doc, ["ความสามารถ", "พฤติกรรม", "การควบคุมข้อมูล"], [
    ("แก้ชื่อ", "เปลี่ยนชื่อวัตถุดิบจากหน้าจอ", "ตรวจชื่อซ้ำก่อนบันทึก"),
    ("แก้น้ำหนักต่อถาด", "แก้ portion weight สำหรับการคำนวณ", "ตรวจค่าที่ส่งเข้า API"),
    ("แก้เกณฑ์ขั้นต่ำ", "กำหนดจำนวนถาดขั้นต่ำใน Prep", "ใช้เป็นเงื่อนไขแจ้งเตือนแบบสด"),
    ("เก็บถาวร", "เปลี่ยน is_active เป็น false", "กันเก็บถาวรหากยังใช้ใน BOM ที่ active"),
    ("เปิดใช้งานใหม่", "เปลี่ยน is_active กลับเป็น true", "กลับมาใช้ใน workflow ใหม่ได้"),
], widths=[1.35, 2.65, 2.85])

add_heading(doc, "การเปลี่ยนฐานข้อมูล", 2)
add_body(doc, "เพิ่ม migration supabase/migrations/0012_ingredient_archiving.sql โดยเพิ่ม ingredients.is_active แบบ BOOLEAN NOT NULL DEFAULT true และ index สำหรับ is_active กับ name การนำ migration ไปใช้กับ Supabase รักษาข้อมูลเดิมครบ 10 รายการ และทุกรายการยัง active หลัง migration")

add_heading(doc, "ผลกระทบต่อ API และหน้าจอ", 2)
add_bullet(doc, "GET /owner/ingredients ส่งคืนทั้งรายการ active และ archived เพื่อให้ Owner จัดการได้")
add_bullet(doc, "PUT /owner/ingredients/:id รองรับชื่อ น้ำหนักต่อถาด เกณฑ์ขั้นต่ำ และสถานะ พร้อม audit log")
add_bullet(doc, "endpoint ปฏิบัติงาน /inventory/ingredients และ staff alerts กรองเฉพาะ active")
add_bullet(doc, "การสร้างเมนูใหม่ตรวจว่าวัตถุดิบใน BOM ยัง active")
add_bullet(doc, "Menu Management ตัดส่วน portion preset ออกและชี้ไปที่ Ingredient Settings")
add_bullet(doc, "Category ยังเป็น read only เพราะการเปลี่ยนประเภทกระทบหน่วยและ flow ระหว่าง Freezer กับ Prep")

doc.add_page_break()

add_heading(doc, "5 ไฟล์ที่เปลี่ยนและผลการตรวจสอบ", 1)
add_table(doc, ["พื้นที่", "ไฟล์หลัก", "สาระสำคัญ"], [
    ("API", "apps/api/src/index.ts", "Owner ingredient API การ archive การตรวจ BOM และ prep alerts"),
    ("Customer", "apps/customer/src/lib/api.ts\napps/customer/vite.config.ts", "API URL precedence และ allowed host สำหรับ Demo"),
    ("Internal", "App.tsx  routes.tsx  sidebar.tsx", "เพิ่ม route เมนู และ global alert"),
    ("Owner UI", "pages/owner/IngredientSettings.tsx\npages/owner/MenuManagement.tsx", "หน้าตั้งค่ารวมและย้ายความรับผิดชอบออกจากหน้า Menu"),
    ("Staff UI", "Notifications.tsx\nTransferToThawPrep.tsx\nStaffPrepAlert.tsx", "รายละเอียด alert การ preselect และตัวแจ้งเตือนทุกหน้า"),
    ("Types", "apps/internal/src/types/ingredient.ts", "รองรับสถานะ active"),
    ("Database", "supabase/migrations/0012_ingredient_archiving.sql", "เพิ่ม is_active และ index"),
], widths=[1.1, 2.85, 2.9])

add_heading(doc, "ผลการตรวจสอบ", 2)
add_table(doc, ["รายการ", "ผล"], [
    ("API build", "ผ่าน"),
    ("Customer build", "ผ่าน"),
    ("Internal build", "ผ่าน มีคำเตือน bundle ประมาณ 548 KB"),
    ("API tests", "ผ่าน 15 จาก 15"),
    ("Internal lint", "ผ่าน มีคำเตือน Fast Refresh เดิม 2 จุดใน AuthContext และ InventoryContext"),
    ("Tunnel smoke test", "Customer HTTP 200  API health HTTP 200  และ CORS ถูกต้อง"),
], widths=[2.25, 4.6])

add_body(doc, "สถานะ Git ขณะจัดทำเอกสาร: มีไฟล์แก้ไขที่ยังไม่ commit 11 ไฟล์ และไฟล์ใหม่ 2 ไฟล์ การเปลี่ยนแปลงที่ Git แสดงมีประมาณ 444 บรรทัดเพิ่มและ 38 บรรทัดลบ โดยไม่นับเนื้อหาของไฟล์ใหม่ใน diff stat")

doc.add_page_break()

add_heading(doc, "6 การตรวจสอบไฟล์สภาพแวดล้อมและความพร้อมก่อน Deploy", 1)
add_body(doc, "การมี .env แยกตาม api customer และ internal เป็นรูปแบบปกติของ monorepo เพราะแต่ละแอปใช้ตัวแปรต่างกัน ควรเก็บไฟล์ตัวอย่าง .env.example ใน Git แต่ห้าม commit ค่า secret หรือ URL ฐานข้อมูลจริง")

add_table(doc, ["ระดับ", "ประเด็น", "ข้อเสนอแนะ"], [
    ("P1", "ยังไม่มี NODE_ENV=production ที่ชัดเจน", "กำหนดใน production runtime เพราะมีผลต่อ secure cookie dev tools LAN CORS และ worker"),
    ("P1", ".gitignore ไม่ครอบคลุม .env.staging .env.development .env.test", "ใช้รูปแบบ .env.* และยกเว้น !.env.example"),
    ("P1", "Customer Vite config มี Quick Tunnel hostname แบบ hard coded", "ย้ายเป็น environment specific configuration และลบเมื่อเลิก Demo"),
    ("P2", "Frontend fallback ไป host:3000 หรือ localhost เมื่อไม่มี VITE_API_URL", "ให้ production build ล้มเหลวทันทีเมื่อไม่ได้กำหนดค่า"),
    ("P2", "การโหลดและ validate env ยังไม่รวมศูนย์", "เพิ่ม schema validation และหยุดโปรแกรมเมื่อค่าที่จำเป็นหาย"),
], widths=[0.65, 3.0, 3.2])

add_heading(doc, "หลักการสำหรับ Production", 2)
add_bullet(doc, "เก็บ secrets ในระบบจัดการ secrets ของ hosting หรือ CI CD ไม่เก็บใน repository")
add_bullet(doc, "ตัวแปรที่ขึ้นต้น VITE_ เป็นข้อมูลสาธารณะใน browser bundle ห้ามใส่ secret")
add_bullet(doc, "ควรใช้โดเมนจริงและถ้าเป็นไปได้ให้อยู่ภายใต้ parent domain เดียวกันเพื่อจัดการ cookie และ CORS ง่ายขึ้น")
add_bullet(doc, "Quick Tunnel และ URL ชั่วคราวไม่ควรใช้เป็นค่าถาวรใน production")
add_bullet(doc, "Docker เป็นตัวเลือกเพื่อมาตรฐาน environment แต่ไม่ใช่เงื่อนไขบังคับของ deployment นี้")

add_heading(doc, "สิ่งที่ยังไม่ได้ดำเนินการ", 2)
add_body(doc, "ยังไม่ได้ทำ permanent deployment ยังไม่ได้ปรับ env hardening ตามข้อเสนอแนะทั้งหมด ยังไม่ได้ทำ production E2E และยังไม่ได้ commit push หรือ merge โค้ด เอกสารนี้จึงเป็นจุดส่งต่อให้ทีมตรวจสอบก่อนดำเนินการ Git")

doc.add_page_break()

add_heading(doc, "7 เช็กลิสต์ก่อน Push และ Merge", 1)
checklist = [
    "ตรวจ diff ของทุกไฟล์และยืนยันว่าไม่มี URL ชั่วคราวหรือข้อมูลลับติดเข้า commit",
    "ตรวจ migration 0012 ใน environment เป้าหมายและสำรองข้อมูลตามนโยบายทีม",
    "ทดสอบ Owner แก้ชื่อ น้ำหนัก เกณฑ์ขั้นต่ำ เก็บถาวร และเปิดใช้งานใหม่",
    "ทดสอบเงื่อนไขห้ามเก็บถาวรวัตถุดิบที่ยังอยู่ใน active BOM",
    "ทดสอบ Staff alert ตั้งแต่สต็อกต่ำ กดไปหน้า Transfer เติมครบ และ alert หาย",
    "ทดสอบว่า archived ingredient ไม่ปรากฏใน intake transfer alerts และ menu ใหม่ แต่ประวัติเดิมยังอยู่",
    "รัน API Customer Internal build อีกครั้งหลัง resolve conflict",
    "รัน API tests และ lint อีกครั้ง พร้อมบันทึก warning ที่ยอมรับได้ใน pull request",
    "กำหนด production environment และ secrets บน hosting โดยไม่ใช้ Quick Tunnel URL",
    "ให้ reviewer ตรวจผลกระทบของ index migration API contract และสิทธิ์ Owner Staff",
]
for item in checklist:
    p = doc.add_paragraph()
    style_paragraph(p, after=5, line=1.15)
    set_run_font(p.add_run("☐  "), size=11, color=ACCENT)
    set_run_font(p.add_run(item), size=10.5)

add_heading(doc, "ข้อเสนอแนะการแบ่ง Commit", 2)
add_table(doc, ["ลำดับ", "Commit ที่แนะนำ", "ขอบเขต"], [
    ("1", "feat db ingredient archiving", "migration และ type ที่เกี่ยวข้อง"),
    ("2", "feat api ingredient settings and prep alerts", "endpoint validation archive guard และ alert calculation"),
    ("3", "feat internal ingredient settings and staff alerts", "Owner UI Staff UI routes sidebar และ global alert"),
    ("4", "chore demo tunnel configuration", "Customer API URL และ allowedHosts โดยพิจารณาไม่ merge ค่า URL ชั่วคราว"),
], widths=[0.65, 2.65, 3.55])

add_body(doc, "ข้อสรุป: ฟีเจอร์หลักพร้อมสำหรับ code review และการทดสอบร่วมกันในทีม ส่วนการ deploy production ควรทำหลังจากปิดประเด็น environment configuration URL ชั่วคราว security settings และ end to end test แล้ว")

# Keep body fonts deterministic, including text created by list styles.
for paragraph in doc.paragraphs:
    for run in paragraph.runs:
        if run.font.name is None:
            set_run_font(run)

doc.core_properties.title = "สรุปการพัฒนาและปรับปรุงระบบ RIMS"
doc.core_properties.subject = "เอกสารส่งต่องานก่อน Push และ Merge เข้า Git"
doc.core_properties.author = "RIMS Development Team"
doc.core_properties.keywords = "RIMS Git merge handoff inventory prep alert ingredient settings"
doc.save(OUTPUT)
print(OUTPUT)
